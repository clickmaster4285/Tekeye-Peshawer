"""
Agent reports: compiled from the same RBAC/location-scoped datasets as search_records, stored as
AgentReport, and rendered on demand as PDF (reportlab) or Excel (openpyxl).

The figures in a report always come from the database, never from the model; the model only
chooses what to include and may add a narrative summary.
"""

from __future__ import annotations

import io
import os
import re

from django.utils import timezone

from .models import AgentReport
from .records import MAX_REPORT_ROWS, RecordsError, build_query, clean_rows, group_counts

MAX_SECTIONS = 8
REPORT_GROUPS = 25
DEFAULT_REPORT_ROWS = 100


class ReportError(Exception):
    pass


def _label(col: str) -> str:
    return col.replace("__", " ").replace("_", " ").strip().capitalize()


def _criteria(spec: dict, notes: list[str]) -> str:
    parts = []
    if spec.get("text"):
        parts.append(f"matching '{spec['text']}'")
    for key, value in (spec.get("filters") or {}).items():
        parts.append(f"{_label(key).lower()} = {value}")
    if spec.get("date_from") or spec.get("date_to"):
        parts.append(f"from {spec.get('date_from') or 'start'} to {spec.get('date_to') or 'now'}")
    elif spec.get("last_days"):
        parts.append(f"last {spec['last_days']} days")
    text = ", ".join(parts) or "all records"
    return text + (" — " + " ".join(notes) if notes else "")


def _section(user, spec: dict) -> dict:
    if not isinstance(spec, dict) or not spec.get("dataset"):
        raise ReportError("Each section needs a dataset.")
    try:
        rq = build_query(user, spec)
    except RecordsError as exc:
        raise ReportError(str(exc)) from exc
    notes = list(rq.notes)
    if rq.fuzzy:
        notes.append(f"No exact match for '{spec.get('text')}'; closest spellings included.")
    total = rq.qs.count()
    group = str(spec.get("group_by") or "").strip()
    try:
        counts = group_counts(rq, group, REPORT_GROUPS) if group else []
        limit = max(0, min(int(spec.get("limit", DEFAULT_REPORT_ROWS)), MAX_REPORT_ROWS))
    except RecordsError as exc:
        raise ReportError(str(exc)) from exc
    except (TypeError, ValueError):
        limit = DEFAULT_REPORT_ROWS
    rows = clean_rows(rq, limit) if limit else []
    columns = [c for c in rq.ds.fields if c != "id" and any(c in r for r in rows)] if rows else []
    return {
        "heading": str(spec.get("heading") or _label(rq.name)).strip()[:150],
        "dataset": rq.name,
        "description": rq.ds.about,
        "criteria": _criteria(spec, notes),
        "total": total,
        "counts_by": _label(group) if group else "",
        "counts": counts,
        "columns": [{"key": c, "label": _label(c)} for c in columns],
        "rows": [{c: r.get(c, "") for c in columns} for r in rows],
        "truncated": total > len(rows),
    }


def _key_figures(sections: list[dict]) -> list[str]:
    figures = []
    for s in sections:
        line = f"{s['heading']}: {s['total']} record(s)"
        if s["counts"]:
            top = ", ".join(f"{c['value']} {c['count']}" for c in s["counts"][:4])
            line += f" — by {s['counts_by'].lower()}: {top}"
        figures.append(line)
    return figures


def build_report(user, session, args: dict) -> AgentReport:
    title = str(args.get("title") or "").strip()[:200]
    if not title:
        raise ReportError("A report title is required.")
    specs = args.get("sections")
    if not specs and args.get("dataset"):
        specs = [{k: v for k, v in args.items() if k not in ("title", "summary", "sections")}]
    if not isinstance(specs, list) or not specs:
        raise ReportError("Give at least one section: {dataset, text?, filters?, group_by?, date_from?, date_to?, last_days?}.")
    sections = [_section(user, spec) for spec in specs[:MAX_SECTIONS]]
    summary = str(args.get("summary") or "").strip()[:6000]
    return AgentReport.objects.create(
        user=user,
        session=session,
        title=title,
        summary=summary,
        sections=sections,
    )


def report_digest(report: AgentReport) -> dict:
    """What the model sees after generating: the real figures, so its narrative matches the document."""
    return {
        "report_id": str(report.pk),
        "title": report.title,
        "key_figures": _key_figures(report.sections),
        "sections": [
            {"heading": s["heading"], "total": s["total"], "counts": s["counts"][:10], "criteria": s["criteria"],
             "rows_included": len(s["rows"])}
            for s in report.sections
        ],
        "instruction": "The report is on the officer's screen with PDF and Excel download buttons. "
                       "Summarise the key findings from these figures in your reply.",
    }


def report_payload(report: AgentReport) -> dict:
    return {
        "id": str(report.pk),
        "title": report.title,
        "summary": report.summary,
        "key_figures": _key_figures(report.sections),
        "sections": report.sections,
        "created_at": timezone.localtime(report.created_at).strftime("%Y-%m-%d %H:%M"),
        "created_by": getattr(report.user, "full_name", "") or report.user.username,
    }


def report_filename(report: AgentReport, ext: str) -> str:
    safe = re.sub(r"[^\w\-]+", "-", report.title).strip("-")[:60] or "report"
    return f"{safe}-{timezone.localtime(report.created_at):%Y%m%d-%H%M}.{ext}"


# ---------------------------------------------------------------- PDF

_FONT_CANDIDATES = (
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("Arial", r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf"),
)


def _fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name, regular, bold in _FONT_CANDIDATES:
        if os.path.exists(regular) and os.path.exists(bold):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, regular))
                pdfmetrics.registerFont(TTFont(f"{name}-Bold", bold))
            return name, f"{name}-Bold"
    return "Helvetica", "Helvetica-Bold"


def _bar_chart(counts: list[dict], width: float, font: str):
    from reportlab.graphics.charts.barcharts import HorizontalBarChart
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors

    data = counts[:12]
    bar_h = 14
    height = max(60, bar_h * len(data) + 30)
    drawing = Drawing(width, height)
    chart = HorizontalBarChart()
    chart.x, chart.y = 150, 15
    chart.width, chart.height = width - 180, height - 25
    chart.data = [[c["count"] for c in reversed(data)]]
    chart.categoryAxis.categoryNames = [str(c["value"])[:28] for c in reversed(data)]
    chart.categoryAxis.labels.fontName = font
    chart.categoryAxis.labels.fontSize = 7
    chart.categoryAxis.labels.boxAnchor = "e"
    chart.valueAxis.labels.fontName = font
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.bars[0].fillColor = colors.HexColor("#2563eb")
    chart.bars[0].strokeColor = None
    chart.barLabels.fontName = font
    chart.barLabels.fontSize = 6.5
    chart.barLabelFormat = "%d"
    chart.barLabels.nudge = 8
    drawing.add(chart)
    return drawing


def render_pdf(report: AgentReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from xml.sax.saxutils import escape

    font, bold = _fonts()
    wide = any(len(s["columns"]) > 6 for s in report.sections)
    pagesize = landscape(A4) if wide else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=pagesize, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=16 * mm, bottomMargin=14 * mm,
        title=report.title, author="TekEye",
    )
    width = pagesize[0] - 28 * mm
    ink, muted, accent = colors.HexColor("#0f172a"), colors.HexColor("#64748b"), colors.HexColor("#1e3a8a")
    st = {
        "brand": ParagraphStyle("brand", fontName=bold, fontSize=8, textColor=accent, spaceAfter=2),
        "title": ParagraphStyle("title", fontName=bold, fontSize=17, leading=21, textColor=ink, spaceAfter=3),
        "meta": ParagraphStyle("meta", fontName=font, fontSize=8, textColor=muted, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName=bold, fontSize=12, leading=15, textColor=ink, spaceBefore=10, spaceAfter=3),
        "body": ParagraphStyle("body", fontName=font, fontSize=9, leading=12.5, textColor=ink, spaceAfter=4),
        "small": ParagraphStyle("small", fontName=font, fontSize=7.5, leading=9.5, textColor=muted, spaceAfter=4),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=7, leading=8.5, textColor=ink),
        "head": ParagraphStyle("head", fontName=bold, fontSize=7, leading=8.5, textColor=colors.white),
    }
    p = lambda text, style: Paragraph(escape(str(text)), st[style])  # noqa: E731
    payload = report_payload(report)

    story = [
        p("PAKISTAN CUSTOMS — TEKEYE", "brand"),
        p(report.title, "title"),
        p(f"Generated {payload['created_at']} by {payload['created_by']} · {len(report.sections)} section(s)", "meta"),
    ]
    if report.summary:
        story.append(p("Summary", "h2"))
        for para in [x for x in re.split(r"\n{2,}", report.summary) if x.strip()]:
            story.append(p(re.sub(r"[*_`#]", "", para.strip()), "body"))
    story.append(p("Key figures", "h2"))
    for line in payload["key_figures"]:
        story.append(p(f"• {line}", "body"))

    for section in report.sections:
        story.append(p(section["heading"], "h2"))
        story.append(p(f"{section['total']} record(s) — {section['criteria']}", "small"))
        if section["counts"]:
            story.append(p(f"By {section['counts_by'].lower()}", "small"))
            story.append(_bar_chart(section["counts"], width, font))
            story.append(Spacer(1, 4))
        if section["rows"]:
            cols = section["columns"]
            data = [[p(c["label"], "head") for c in cols]]
            data += [[p(r.get(c["key"], ""), "cell") for c in cols] for r in section["rows"]]
            table = LongTable(data, colWidths=[width / len(cols)] * len(cols), repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), accent),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                    ]
                )
            )
            story.append(table)
            if section["truncated"]:
                story.append(p(f"Showing {len(section['rows'])} of {section['total']} records.", "small"))
        elif not section["total"]:
            story.append(p("No records matched.", "body"))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(muted)
        canvas.drawString(14 * mm, 8 * mm, f"{report.title} — TekEye agent report")
        canvas.drawRightString(pagesize[0] - 14 * mm, 8 * mm, f"Page {_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel


def render_xlsx(report: AgentReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    payload = report_payload(report)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated {payload['created_at']} by {payload['created_by']}"
    row = 4
    if report.summary:
        ws.cell(row=row, column=1, value=re.sub(r"[*_`#]", "", report.summary)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = min(400, 15 * (1 + report.summary.count("\n") + len(report.summary) // 110))
        row += 2
    ws.cell(row=row, column=1, value="Key figures").font = Font(bold=True)
    for line in payload["key_figures"]:
        row += 1
        ws.cell(row=row, column=1, value=line)
    ws.column_dimensions["A"].width = 120

    used = {"Summary"}
    for i, section in enumerate(report.sections, start=1):
        name = re.sub(r"[\[\]:*?/\\]", "", section["heading"])[:28] or f"Section {i}"
        while name in used:
            name = f"{name[:25]}-{i}"
        used.add(name)
        sh = wb.create_sheet(name)
        sh["A1"] = section["heading"]
        sh["A1"].font = Font(bold=True, size=12)
        sh["A2"] = f"{section['total']} record(s) — {section['criteria']}"
        r = 4
        if section["counts"]:
            sh.cell(row=r, column=1, value=section["counts_by"] or "Value")
            sh.cell(row=r, column=2, value="Count")
            for c in (1, 2):
                sh.cell(row=r, column=c).font, sh.cell(row=r, column=c).fill = head_font, head_fill
            for item in section["counts"]:
                r += 1
                sh.cell(row=r, column=1, value=str(item["value"]))
                sh.cell(row=r, column=2, value=item["count"])
            chart = BarChart()
            chart.type = "bar"
            chart.title = f"By {section['counts_by'].lower()}"
            chart.legend = None
            chart.add_data(Reference(sh, min_col=2, min_row=4, max_row=r), titles_from_data=True)
            chart.set_categories(Reference(sh, min_col=1, min_row=5, max_row=r))
            chart.height = max(6, 0.6 * len(section["counts"]) + 2)
            chart.width = 16
            sh.add_chart(chart, "D4")
            r = max(r, 4 + int(chart.height * 2)) + 2
        if section["rows"]:
            start = r
            for j, col in enumerate(section["columns"], start=1):
                cell = sh.cell(row=r, column=j, value=col["label"])
                cell.font, cell.fill = head_font, head_fill
            for data_row in section["rows"]:
                r += 1
                for j, col in enumerate(section["columns"], start=1):
                    sh.cell(row=r, column=j, value=data_row.get(col["key"], ""))
            last = get_column_letter(len(section["columns"]))
            sh.auto_filter.ref = f"A{start}:{last}{r}"
            sh.freeze_panes = sh.cell(row=start + 1, column=1)
            for j, col in enumerate(section["columns"], start=1):
                longest = max([len(col["label"])] + [len(str(x.get(col["key"], ""))) for x in section["rows"][:200]])
                sh.column_dimensions[get_column_letter(j)].width = min(50, max(10, longest + 2))
            if section["truncated"]:
                sh.cell(row=r + 2, column=1, value=f"Showing {len(section['rows'])} of {section['total']} records.")
        elif not section["total"]:
            sh.cell(row=r, column=1, value="No records matched.")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
